"""Exportar la misma instantánea consultada, con texto seguro y horas numéricas."""
import os
import tempfile
from datetime import datetime, timedelta
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from .formatting import duration_text, seconds_from_hours


def append_safe(sheet, values):
    sheet.append(values)
    for cell in sheet[sheet.max_row]:
        # Evita interpretar nombres/áreas iniciados en '=' como fórmulas.
        if isinstance(cell.value, str):
            cell.data_type = 's'


def style_table(sheet):
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = PatternFill('solid', fgColor='B51C67')
        cell.font = Font(size=14, bold=True, color='FFFFFF')
        cell.alignment = Alignment(vertical='center')
    sheet.row_dimensions[1].height = 36
    for column in sheet.columns:
        letter = column[0].column_letter
        sheet.column_dimensions[letter].width = min(48, max(16, max(len(str(c.value or '')) for c in column)+2))
    for row in sheet.iter_rows(min_row=2):
        sheet.row_dimensions[row[0].row].height = 27
        for cell in row:
            cell.font = Font(size=13)
            if isinstance(cell.value, datetime):
                cell.number_format = 'dd/mm/yyyy hh:mm:ss'
            elif isinstance(cell.value, float):
                cell.number_format = '0.00'


def write_duration(cell, seconds):
    """Duración numérica sumable en Excel, presentada sin horas decimales."""
    if seconds is None:
        cell.value = None
        return
    seconds = max(0, int(seconds))
    cell.value = timedelta(seconds=seconds)
    if seconds < 60:
        cell.number_format = '"0 minutos"' if seconds == 0 else '"Menos de 1 minuto"'
    elif seconds < 3600:
        cell.number_format = '[m] "minuto"' if seconds < 120 else '[m] "minutos"'
    else:
        hours, minutes = divmod(seconds // 60, 60)
        cell.number_format = f'[h] "{"hora" if hours == 1 else "horas"}" m "{"minuto" if minutes == 1 else "minutos"}"'


def export_report(report, destination):
    destination = Path(destination)
    wb = Workbook()
    detail = wb.active
    detail.title = 'Detalle'
    append_safe(detail, ['Registro', 'Documento', 'Apellidos y nombres', 'Tipo', 'Cargo o área',
                         'Ingreso real', 'Salida real', 'Tiempo sesión completa', 'Tiempo en periodo', 'Estado'])
    for row in report.details:
        append_safe(detail, [row['id'], row['numero_documento'], row['nombre'], row['tipo_persona'], row['cargo_o_area'],
            datetime.fromisoformat(row['fecha_ingreso']), datetime.fromisoformat(row['fecha_salida']) if row['fecha_salida'] else None,
            row['horas_trabajadas'], row['hours_period'], 'Cerrado' if row['fecha_salida'] else 'Pendiente'])
    summary = wb.create_sheet('Resumen por persona')
    append_safe(summary, ['Documento', 'Apellidos y nombres', 'Tipo', 'Cargo o área', 'Días con presencia cerrada', 'Tiempo en periodo', 'Ingresos pendientes'])
    for row in report.summary:
        append_safe(summary, [row['numero_documento'], row['nombre'], row['tipo_persona'], row['cargo_o_area'], row['days'], row['hours'], row['pending']])
    meta = wb.create_sheet('Criterios')
    for values in [
        ['Criterio', 'Valor'], ['Desde (incluido)', report.start], ['Hasta (incluido)', report.end-timedelta(seconds=1)],
        ['Búsqueda', report.query or 'Todas las personas'], ['Tipo', report.kind], ['Generado', report.generated_at],
        ['Lugar y hora local', 'Bogotá/Colombia'], ['Tiempo total confirmado', duration_text(sum(r['seconds'] for r in report.details))],
        ['Pendientes', report.open_count], ['Días', 'Fechas distintas con presencia cerrada, por persona; incluye fines de semana.'],
        ['Tiempo', 'Duraciones numéricas en Excel, mostradas en minutos u horas y minutos; no equivale a nómina.'],
        ['Pendientes', 'Sin horas hasta registrar o corregir la salida.'],
        ['Identidad', 'Los datos de persona corresponden a su ficha actual.'],
    ]:
        append_safe(meta, values)
    for sheet in wb:
        style_table(sheet)
    for row in detail.iter_rows(min_row=2):
        row[1].number_format = '@'
    for row in summary.iter_rows(min_row=2):
        row[0].number_format = '@'
    for index, row in enumerate(report.details, start=2):
        write_duration(detail.cell(index, 8), seconds_from_hours(row['horas_trabajadas']))
        write_duration(detail.cell(index, 9), row['seconds'] if row['fecha_salida'] else None)
    for index, row in enumerate(report.summary, start=2):
        write_duration(summary.cell(index, 6), row['seconds'])
    # Guardado atómico: si Excel mantiene el destino abierto, no corromperlo.
    handle, temporary = tempfile.mkstemp(suffix='.xlsx', dir=destination.parent)
    os.close(handle)
    try:
        wb.save(temporary)
        os.replace(temporary, destination)
    finally:
        wb.close()
        Path(temporary).unlink(missing_ok=True)
